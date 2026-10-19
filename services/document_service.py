"""DOCX : reconstruction en Word avec python-docx 1.2 et quelques fragments OOXML isolés.

Principes (TECHNICAL_DECISIONS D-05, D-06, D-07) :
- une section « nouvelle page » par page source ; les bandes de colonnes supplémentaires d'une page sont des
  sections continues (`w:cols`) ; le changement de section fait le saut de page (aucun saut explicite) ;
- styles Word réutilisables créés une fois, valeurs de base = valeur la plus fréquente du livre pour ce type
  de bloc ; chaque paragraphe ne porte en direct que ses écarts au style ;
- texte "absolute" -> paragraphe dans un cadre `w:framePr` (hRule « atLeast » : le cadre grandit au lieu de
  couper le texte ; un texte édité plus long peut chevaucher un voisin, jamais disparaître) ;
- image "absolute" et fonds -> `wp:anchor` (fond : derrière le texte, sans habillage) ; image "flow" -> inline ;
- "header" / "footer" -> en-tête / pied de la section de la page ;
- le texte n'est jamais remplacé par une image ; tout repli est listé dans `ExportReport.fallbacks`.
"""
from __future__ import annotations

import copy
import logging
import os
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Literal

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor
from lxml import etree
from pydantic import BaseModel, Field

from models.blocks import Alignment, BlockKind
from models.composition import ColumnBand, ComposedPage, ImageElement, Rect, ResolvedStyle, TextElement
from models.errors import ExportError
from models.project import Project, ReconstructionMode
from models.reports import ExportReport, Issue, Severity
from services import fonts

logger = logging.getLogger("bookocr.docx")

TWIPS_PER_PT = 20
EMU_PER_PT = 12700
FULL_PAGE_RATIO = 0.9

STYLE_NAMES: dict[BlockKind, str] = {
    BlockKind.HEADING_1: "Titre 1",
    BlockKind.HEADING_2: "Titre 2",
    BlockKind.PARAGRAPH: "Corps",
    BlockKind.LIST_ITEM: "Liste",
    BlockKind.CAPTION: "Légende",
    BlockKind.HEADER: "En-tête livre",
    BlockKind.FOOTER: "Pied livre",
    BlockKind.PAGE_NUMBER: "Numéro de page",
    BlockKind.NOTE: "Note",
    BlockKind.TABLE: "Tableau texte",
    BlockKind.SPECIAL: "Spécial",
}
OUTLINE_LEVELS = {BlockKind.HEADING_1: 0, BlockKind.HEADING_2: 1}
IMAGE_PARAGRAPH_STYLE = "Normal"

# Valeurs de base d'un style quand le livre n'a aucun bloc de ce type : (taille, gras, italique, alignement)
DEFAULT_KIND_LOOK: dict[BlockKind, tuple[float, bool, bool, Alignment]] = {
    BlockKind.HEADING_1: (20.0, True, False, Alignment.CENTER),
    BlockKind.HEADING_2: (14.0, True, False, Alignment.LEFT),
    BlockKind.PARAGRAPH: (11.0, False, False, Alignment.JUSTIFY),
    BlockKind.LIST_ITEM: (11.0, False, False, Alignment.LEFT),
    BlockKind.CAPTION: (9.0, False, True, Alignment.CENTER),
    BlockKind.HEADER: (9.0, False, True, Alignment.CENTER),
    BlockKind.FOOTER: (9.0, False, False, Alignment.CENTER),
    BlockKind.PAGE_NUMBER: (9.0, False, False, Alignment.CENTER),
    BlockKind.NOTE: (9.0, False, False, Alignment.JUSTIFY),
    BlockKind.TABLE: (10.0, False, False, Alignment.LEFT),
    BlockKind.SPECIAL: (11.0, False, True, Alignment.CENTER),
}

WORD_ALIGN = {Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT, Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
              Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT, Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY}

# Tolérances sous lesquelles un écart au style n'est pas écrit en direct (bruit d'estimation)
SIZE_TOLERANCE_PT = 0.25
SPACING_TOLERANCE_PT = 0.75
LINE_SPACING_TOLERANCE = 0.02

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

# Ordre des enfants de w:pPr et de w:sectPr (ECMA-376, CT_PPrBase et CT_SectPr)
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]

ANCHOR_REQUIRED_ATTRIBUTES = ["distT", "distB", "distL", "distR", "simplePos", "relativeHeight", "behindDoc",
                              "locked", "layoutInCell", "allowOverlap"]
ANCHOR_CHILD_ORDER = ["simplePos", "positionH", "positionV", "extent", "effectExtent", "wrapNone", "wrapSquare",
                      "wrapTight", "wrapThrough", "wrapTopAndBottom", "docPr", "cNvGraphicFramePr", "graphic"]
ANCHOR_REQUIRED_CHILDREN = ["simplePos", "positionH", "positionV", "extent", "docPr", "graphic"]
WRAP_ELEMENTS = {"wrapNone", "wrapSquare", "wrapTight", "wrapThrough", "wrapTopAndBottom"}

FrameWrap = Literal["none", "around", "notBeside"]
AnchorWrap = Literal["none", "square"]


# --------------------------------------------------------------------------- conversions et XML isolé


def pt_to_twips(value: float) -> int:
    return int(round(value * TWIPS_PER_PT))


def pt_to_emu(value: float) -> int:
    return int(round(value * EMU_PER_PT))


def insert_in_order(parent: etree._Element, child: etree._Element, order: list[str]) -> etree._Element:
    """Insère `child` dans `parent` en respectant l'ordre de schéma `order` (noms locaux w:)."""
    name = etree.QName(child).localname
    existing = parent.find(qn(f"w:{name}"))
    if existing is not None:
        parent.replace(existing, child)
        return child
    successors = set(order[order.index(name) + 1:])
    for index, sibling in enumerate(parent):
        if etree.QName(sibling).localname in successors:
            parent.insert(index, child)
            return child
    parent.append(child)
    return child


def frame_element(rect: Rect, wrap: FrameWrap = "none",
                  height_rule: Literal["atLeast", "exact"] = "atLeast") -> etree._Element:
    """`w:framePr` positionnant un paragraphe sur la page (x, y, largeur, hauteur en twips depuis le coin haut gauche).

    hRule « atLeast » par défaut : la hauteur détectée est un minimum ; un texte plus long agrandit le cadre
    au lieu d'être coupé (avec « exact », Word masque le surplus sans prévenir).
    """
    return parse_xml(
        f'<w:framePr {nsdecls("w")} w:w="{pt_to_twips(rect.width_pt)}" w:h="{pt_to_twips(rect.height_pt)}" '
        f'w:hRule="{height_rule}" w:wrap="{wrap}" w:hAnchor="page" w:vAnchor="page" '
        f'w:x="{pt_to_twips(rect.x_pt)}" w:y="{pt_to_twips(rect.y_pt)}" w:hSpace="0" w:vSpace="0"/>')


def set_frame(paragraph, rect: Rect, wrap: FrameWrap = "none") -> etree._Element:
    p_pr = paragraph._p.get_or_add_pPr()
    return insert_in_order(p_pr, frame_element(rect, wrap), PPR_ORDER)


def columns_element(count: int, gap_pt: float, widths_pt: list[float] | None = None) -> etree._Element:
    """`w:cols` : `count` colonnes séparées de `gap_pt` ; largeurs inégales si `widths_pt` est donné."""
    if widths_pt and len(widths_pt) == count and count > 1:
        cols = "".join(f'<w:col w:w="{pt_to_twips(width)}" w:space="{pt_to_twips(gap_pt)}"/>'
                       for width in widths_pt)
        return parse_xml(f'<w:cols {nsdecls("w")} w:num="{count}" w:space="{pt_to_twips(gap_pt)}" '
                         f'w:equalWidth="0">{cols}</w:cols>')
    if count <= 1:   # une colonne : w:num omis (valeur par défaut), comme Word
        return parse_xml(f'<w:cols {nsdecls("w")} w:space="{pt_to_twips(gap_pt)}"/>')
    return parse_xml(f'<w:cols {nsdecls("w")} w:num="{count}" w:space="{pt_to_twips(gap_pt)}"/>')


def set_section_columns(section, count: int, gap_pt: float, widths_pt: list[float] | None = None) -> None:
    insert_in_order(section._sectPr, columns_element(count, gap_pt, widths_pt), SECTPR_ORDER)


def anchor_from_inline(inline: etree._Element, rect: Rect, *, behind: bool, wrap: AnchorWrap,
                       relative_height: int, distance_pt: float = 0.0) -> etree._Element:
    """Construit un `wp:anchor` positionné sur la page à partir du `wp:inline` créé par python-docx.

    Tous les attributs et enfants obligatoires sont présents (Word refuse une ancre incomplète).
    """
    distance = pt_to_emu(distance_pt)
    wrap_xml = '<wp:wrapNone/>' if wrap == "none" else '<wp:wrapSquare wrapText="bothSides"/>'
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="{distance}" distB="{distance}" '
        f'distL="{distance}" distR="{distance}" simplePos="0" relativeHeight="{relative_height}" '
        f'behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{pt_to_emu(rect.x_pt)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{pt_to_emu(rect.y_pt)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{pt_to_emu(rect.width_pt)}" cy="{pt_to_emu(rect.height_pt)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'{wrap_xml}'
        '</wp:anchor>')
    for name in ("docPr", "cNvGraphicFramePr"):
        found = inline.find(f"wp:{name}", NS)
        if found is not None:
            anchor.append(copy.deepcopy(found))
    graphic = inline.find("a:graphic", NS)
    if graphic is None:
        raise ExportError("Image DOCX invalide (élément a:graphic absent).")
    anchor.append(copy.deepcopy(graphic))
    # l'étendue de l'image elle-même (pic:spPr/a:xfrm/a:ext) suit le rect
    for ext in anchor.iter(qn("a:ext")):
        if ext.get("cx") is None:          # a:ext d'une liste d'extensions (a:extLst), pas une taille
            continue
        ext.set("cx", str(pt_to_emu(rect.width_pt)))
        ext.set("cy", str(pt_to_emu(rect.height_pt)))
    return anchor


def anchor_problems(anchor: etree._Element) -> list[str]:
    """Contrôle de schéma minimal d'un `wp:anchor` : attributs et enfants obligatoires, ordre des enfants."""
    problems = [f"attribut {name} absent" for name in ANCHOR_REQUIRED_ATTRIBUTES if anchor.get(name) is None]
    children = [etree.QName(child).localname for child in anchor]
    problems += [f"enfant {name} absent" for name in ANCHOR_REQUIRED_CHILDREN if name not in children]
    if not WRAP_ELEMENTS & set(children):
        problems.append("aucun élément d'habillage (wrap*)")
    known = [name for name in children if name in ANCHOR_CHILD_ORDER]
    ranks = [ANCHOR_CHILD_ORDER.index(name) for name in known]
    if ranks != sorted(ranks):
        problems.append(f"ordre des enfants invalide : {known}")
    for axis in ("positionH", "positionV"):
        element = anchor.find(f"wp:{axis}", NS)
        if element is not None and element.get("relativeFrom") is None:
            problems.append(f"{axis} sans relativeFrom")
    return problems


def set_run_fonts(r_pr: etree._Element, family: str) -> None:
    """Nom de police Word pour toutes les écritures (ascii, hAnsi, eastAsia, cs)."""
    r_fonts = r_pr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), family)


def set_outline_level(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    insert_in_order(p_pr, parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level}"/>'), PPR_ORDER)


# --------------------------------------------------------------------------- styles


def effective_style(element: TextElement) -> ResolvedStyle:
    """Style tel qu'écrit dans Word : en flux, l'espace avant est l'écart d'origine (gap_before_pt) et l'espace
    après est nul (Word additionne les deux) ; dans un cadre, en-tête ou pied, aucun espacement."""
    if element.positioning == "flow":
        return element.style.model_copy(update={"space_before_pt": round(element.gap_before_pt * 2) / 2,
                                                "space_after_pt": 0.0})
    return element.style.model_copy(update={"space_before_pt": 0.0, "space_after_pt": 0.0})


def _most_common(values: list):
    return Counter(values).most_common(1)[0][0]


def style_baselines(pages: list[ComposedPage], project: Project) -> dict[BlockKind, ResolvedStyle]:
    """Valeurs de base de chaque style : attribut par attribut, la valeur la plus fréquente du livre."""
    by_kind: dict[BlockKind, list[ResolvedStyle]] = {}
    for page in pages:
        for element in page.text_elements:
            by_kind.setdefault(element.kind, []).append(effective_style(element))
    body_sizes = [style.font_size_pt for style in by_kind.get(BlockKind.PARAGRAPH, [])]
    body_size = _most_common(body_sizes) if body_sizes else DEFAULT_KIND_LOOK[BlockKind.PARAGRAPH][0]
    baselines = {}
    for kind in BlockKind:
        styles = by_kind.get(kind)
        if styles:
            fields = ResolvedStyle.model_fields
            baselines[kind] = ResolvedStyle(**{name: _most_common([getattr(style, name) for style in styles])
                                               for name in fields})
            continue
        size, bold, italic, alignment = DEFAULT_KIND_LOOK[kind]
        if kind in (BlockKind.PARAGRAPH, BlockKind.LIST_ITEM):
            size = body_size
        baselines[kind] = ResolvedStyle(font_family=project.typography.document_font, font_size_pt=size,
                                        bold=bold, italic=italic, alignment=alignment,
                                        color=project.typography.default_color,
                                        line_spacing=project.typography.default_line_spacing)
    return baselines


def _apply_font(font, r_pr, style: ResolvedStyle) -> None:
    font.name = style.font_family
    set_run_fonts(r_pr, style.font_family)
    font.size = Pt(style.font_size_pt)
    font.bold = style.bold
    font.italic = style.italic
    font.underline = style.underline
    font.color.rgb = RGBColor.from_string(style.color.lstrip("#").upper())


def _apply_paragraph_format(paragraph_format, style: ResolvedStyle) -> None:
    paragraph_format.alignment = WORD_ALIGN[style.alignment]
    paragraph_format.line_spacing = style.line_spacing
    paragraph_format.space_before = Pt(style.space_before_pt)
    paragraph_format.space_after = Pt(style.space_after_pt)
    paragraph_format.first_line_indent = Pt(style.first_line_indent_pt)
    paragraph_format.left_indent = Pt(style.left_indent_pt)
    paragraph_format.right_indent = Pt(style.right_indent_pt)


def create_styles(document, baselines: dict[BlockKind, ResolvedStyle]) -> dict[BlockKind, str]:
    """Crée une fois les styles de paragraphe du livre ; renvoie type de bloc -> nom de style."""
    normal = document.styles["Normal"]
    body = baselines[BlockKind.PARAGRAPH]
    normal.font.name = body.font_family
    set_run_fonts(normal.element.get_or_add_rPr(), body.font_family)
    normal.font.size = Pt(body.font_size_pt)
    normal.paragraph_format.space_after = Pt(0)
    existing = {style.name for style in document.styles}
    for kind, name in STYLE_NAMES.items():
        style = document.styles[name] if name in existing else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.quick_style = True
        _apply_font(style.font, style.element.get_or_add_rPr(), baselines[kind])
        _apply_paragraph_format(style.paragraph_format, baselines[kind])
        if kind in OUTLINE_LEVELS:
            set_outline_level(style, OUTLINE_LEVELS[kind])
            style.paragraph_format.keep_with_next = True
    for kind in OUTLINE_LEVELS:
        document.styles[STYLE_NAMES[kind]].next_paragraph_style = document.styles[STYLE_NAMES[BlockKind.PARAGRAPH]]
    return dict(STYLE_NAMES)


def style_differences(style: ResolvedStyle, base: ResolvedStyle) -> dict[str, object]:
    """Attributs de `style` qui s'écartent de `base` au-delà des tolérances (à écrire en direct)."""
    differences: dict[str, object] = {}
    for name in ("font_family", "bold", "italic", "underline", "alignment"):
        if getattr(style, name) != getattr(base, name):
            differences[name] = getattr(style, name)
    if style.color.upper() != base.color.upper():
        differences["color"] = style.color
    if abs(style.font_size_pt - base.font_size_pt) > SIZE_TOLERANCE_PT:
        differences["font_size_pt"] = style.font_size_pt
    if abs(style.line_spacing - base.line_spacing) > LINE_SPACING_TOLERANCE:
        differences["line_spacing"] = style.line_spacing
    for name in ("space_before_pt", "space_after_pt", "first_line_indent_pt", "left_indent_pt", "right_indent_pt"):
        if abs(getattr(style, name) - getattr(base, name)) > SPACING_TOLERANCE_PT:
            differences[name] = getattr(style, name)
    return differences


def apply_direct_formatting(paragraph, run, differences: dict[str, object]) -> None:
    """Écrit en direct sur le paragraphe et son texte les seuls écarts au style."""
    paragraph_format = paragraph.paragraph_format
    font = run.font
    for name, value in differences.items():
        if name == "font_family":
            font.name = value
            set_run_fonts(run._r.get_or_add_rPr(), str(value))
        elif name == "font_size_pt":
            font.size = Pt(value)
        elif name in ("bold", "italic", "underline"):
            setattr(font, name, value)
        elif name == "color":
            font.color.rgb = RGBColor.from_string(str(value).lstrip("#").upper())
        elif name == "alignment":
            paragraph_format.alignment = WORD_ALIGN[value]
        elif name == "line_spacing":
            paragraph_format.line_spacing = value
        elif name == "space_before_pt":
            paragraph_format.space_before = Pt(value)
        elif name == "space_after_pt":
            paragraph_format.space_after = Pt(value)
        elif name == "first_line_indent_pt":
            paragraph_format.first_line_indent = Pt(value)
        elif name == "left_indent_pt":
            paragraph_format.left_indent = Pt(value)
        elif name == "right_indent_pt":
            paragraph_format.right_indent = Pt(value)


# --------------------------------------------------------------------------- construction


class _Builder:
    """État de construction d'un document (styles, compteurs, rapport)."""

    def __init__(self, document, project: Project, pages: list[ComposedPage], report: ExportReport) -> None:
        self.document = document
        self.project = project
        self.report = report
        self.baselines = style_baselines(pages, project)
        self.style_names = create_styles(document, self.baselines)
        self.relative_height = 251_658_240
        self.previous_header_signature: tuple = ()
        self.previous_footer_signature: tuple = ()
        self.inserted_images = 0
        self.expected_texts: list[str] = []

    # ---- sections

    def configure_section(self, section, page: ComposedPage, start: WD_SECTION) -> None:
        section.start_type = start
        section.orientation = WD_ORIENT.LANDSCAPE if page.width_pt > page.height_pt else WD_ORIENT.PORTRAIT
        section.page_width = Emu(pt_to_emu(page.width_pt))
        section.page_height = Emu(pt_to_emu(page.height_pt))
        margins = page.margins
        horizontal_room = page.width_pt - 36
        left = min(max(margins.left_pt, 0.0), horizontal_room / 2)
        right = min(max(margins.right_pt, 0.0), horizontal_room / 2)
        top = min(max(margins.top_pt, 0.0), (page.height_pt - 36) / 2)
        bottom = min(max(margins.bottom_pt, 0.0), (page.height_pt - 36) / 2)
        if (left, right, top, bottom) != (margins.left_pt, margins.right_pt, margins.top_pt, margins.bottom_pt):
            self.report.fallbacks.append(Issue(
                code="margins_adjusted", severity=Severity.INFO, page_id=page.page_id,
                message=f"Page {page.page_number} : marges ramenées dans les limites de la page."))
        section.left_margin, section.right_margin = Pt(left), Pt(right)
        section.top_margin, section.bottom_margin = Pt(top), Pt(bottom)
        section.gutter = Pt(0)
        section.header_distance = Pt(min(top, 36.0) / 2)
        section.footer_distance = Pt(min(bottom, 36.0) / 2)
        set_section_columns(section, 1, 0.0)

    def configure_band(self, section, page: ComposedPage, band: ColumnBand | None) -> None:
        if band is None or band.columns <= 1:
            set_section_columns(section, 1, 0.0)
            return
        widths = [rect.width_pt for rect in band.column_rects]
        unequal = len(widths) == band.columns and max(widths) - min(widths) > 3.0
        set_section_columns(section, band.columns, band.gap_pt, widths if unequal else None)

    # ---- texte

    def write_text(self, paragraph, element: TextElement, column_break: bool = False) -> None:
        if column_break:
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
        run = paragraph.add_run(element.text)
        differences = style_differences(effective_style(element), self.baselines[element.kind])
        apply_direct_formatting(paragraph, run, differences)
        self.report.text_blocks += 1
        self.expected_texts.append(element.text)

    def add_text_paragraph(self, element: TextElement, column_break: bool = False):
        paragraph = self.document.add_paragraph(style=self.style_names[element.kind])
        self.write_text(paragraph, element, column_break)
        return paragraph

    def add_frame(self, page: ComposedPage, element: TextElement) -> None:
        paragraph = self.add_text_paragraph(element)
        wrap: FrameWrap = "around" if page.mode == ReconstructionMode.HYBRID else "none"
        set_frame(paragraph, element.rect, wrap)

    # ---- images

    def refuse_image(self, page: ComposedPage, image: ImageElement) -> bool:
        """Vrai (et signalé) si l'image ne doit pas ou ne peut pas être insérée."""
        path = Path(image.file)
        if page.source_image and path == Path(page.source_image):
            self.report.issues.append(Issue(
                code="source_photo_refused", severity=Severity.ERROR, page_id=page.page_id, block_id=image.block_id,
                message=f"Page {page.page_number} : la photo source n'est pas insérée comme image "
                        "(le texte ne doit jamais être remplacé par une image)."))
            return True
        if not path.is_file():
            self.report.issues.append(Issue(
                code="image_missing", severity=Severity.ERROR, page_id=page.page_id, block_id=image.block_id,
                message=f"Page {page.page_number} : image introuvable ({path.name})."))
            return True
        if (page.mode == ReconstructionMode.EDITABLE and image.kind != "background"
                and image.rect.width_pt * image.rect.height_pt >= FULL_PAGE_RATIO * page.width_pt * page.height_pt):
            self.report.issues.append(Issue(
                code="full_page_image_editable", severity=Severity.WARNING, page_id=page.page_id,
                block_id=image.block_id,
                message=f"Page {page.page_number} : image couvrant presque toute la page en mode modifiable."))
        return False

    def add_picture_run(self, paragraph, page: ComposedPage, image: ImageElement, width_pt: float,
                        height_pt: float):
        run = paragraph.add_run()
        try:
            run.add_picture(image.file, width=Emu(pt_to_emu(width_pt)), height=Emu(pt_to_emu(height_pt)))
        except (UnrecognizedImageError, OSError, ValueError) as error:
            self.report.issues.append(Issue(
                code="image_unreadable", severity=Severity.ERROR, page_id=page.page_id, block_id=image.block_id,
                message=f"Page {page.page_number} : image illisible ({Path(image.file).name}) : {error}"))
            paragraph._p.remove(run._r)
            return None
        self.inserted_images += 1
        self.report.images += 1
        descr = "Fond de page" if image.kind == "background" else f"Image {image.block_id or ''}".strip()
        for doc_pr in run._r.iter(qn("wp:docPr")):
            doc_pr.set("descr", descr)
        return run

    def add_anchor(self, carrier, page: ComposedPage, image: ImageElement) -> None:
        if self.refuse_image(page, image):
            return
        run = self.add_picture_run(carrier, page, image, image.rect.width_pt, image.rect.height_pt)
        if run is None:
            return
        inline = run._r.find(".//wp:inline", NS)
        behind = image.kind == "background"
        wrap: AnchorWrap = "none" if behind or page.mode == ReconstructionMode.FIDELITY else "square"
        self.relative_height += 1024
        anchor = anchor_from_inline(inline, image.rect, behind=behind, wrap=wrap,
                                    relative_height=self.relative_height,
                                    distance_pt=0.0 if wrap == "none" else 3.0)
        inline.getparent().replace(inline, anchor)

    def add_inline_image(self, page: ComposedPage, image: ImageElement, column_width_pt: float,
                         column_left_pt: float, column_break: bool) -> None:
        if self.refuse_image(page, image):
            return
        width, height = image.rect.width_pt, image.rect.height_pt
        if width > column_width_pt + 1:
            factor = column_width_pt / width
            width, height = width * factor, height * factor
            self.report.fallbacks.append(Issue(
                code="image_resized", severity=Severity.INFO, page_id=page.page_id, block_id=image.block_id,
                message=f"Page {page.page_number} : image réduite à {factor:.0%} pour tenir dans la colonne."))
        paragraph = self.document.add_paragraph(style=IMAGE_PARAGRAPH_STYLE)
        paragraph_format = paragraph.paragraph_format
        paragraph_format.space_before = Pt(max(image.gap_before_pt, 0.0))
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing = 1.0
        paragraph_format.first_line_indent = Pt(0)
        paragraph_format.keep_with_next = True       # l'image reste avec sa légende
        image_center = image.rect.x_pt + image.rect.width_pt / 2
        column_center = column_left_pt + column_width_pt / 2
        if abs(image_center - column_center) <= 0.1 * column_width_pt:
            paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph_format.left_indent = Pt(max(0.0, image.rect.x_pt - column_left_pt))
        if column_break:
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
        self.add_picture_run(paragraph, page, image, width, height)

    # ---- en-têtes et pieds

    def fill_header_footer(self, section, page: ComposedPage, elements: list[TextElement], footer: bool) -> None:
        signature = tuple((element.kind, element.text, effective_style(element)) for element in elements)
        previous = self.previous_footer_signature if footer else self.previous_header_signature
        part = section.footer if footer else section.header
        if footer:
            self.previous_footer_signature = signature
        else:
            self.previous_header_signature = signature
        if elements:
            if footer:
                section.footer_distance = Pt(max(0.0, page.height_pt - max(e.rect.bottom_pt for e in elements)))
            else:
                section.header_distance = Pt(max(0.0, min(element.rect.y_pt for element in elements)))
        for element in elements:
            self.expected_texts.append(element.text)
        if signature == previous:
            return                              # même contenu que la page précédente : section liée
        part.is_linked_to_previous = False
        paragraphs = list(part.paragraphs)
        for extra in paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        if not elements:
            if paragraphs:
                paragraphs[0].style = self.document.styles["Normal"]
            return
        for index, element in enumerate(elements):
            paragraph = paragraphs[0] if index == 0 and paragraphs else part.add_paragraph()
            paragraph.style = self.document.styles[self.style_names[element.kind]]
            run = paragraph.add_run(element.text)
            apply_direct_formatting(paragraph, run,
                                    style_differences(effective_style(element), self.baselines[element.kind]))
            self.report.text_blocks += 1

    # ---- page

    def add_carrier(self):
        """Paragraphe porteur des images ancrées de la page (interligne exact 1 pt : ne décale pas le flux)."""
        carrier = self.document.add_paragraph(style=IMAGE_PARAGRAPH_STYLE)
        paragraph_format = carrier.paragraph_format
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing = Pt(1)
        return carrier

    def add_page(self, page: ComposedPage, first: bool) -> None:
        section = self.document.sections[0] if first else self.document.add_section(WD_SECTION.NEW_PAGE)
        self.configure_section(section, page, WD_SECTION.NEW_PAGE)
        self.report.issues.extend(page.issues)
        headers = [e for e in page.text_elements if e.positioning == "header" and e.text.strip()]
        footers = [e for e in page.text_elements if e.positioning == "footer" and e.text.strip()]
        if not first or headers:
            self.fill_header_footer(section, page, headers, footer=False)
        if not first or footers:
            self.fill_header_footer(section, page, footers, footer=True)
        if not page.analysed:
            self.document.add_paragraph(style=self.style_names[BlockKind.PARAGRAPH])
            self.report.issues.append(Issue(
                code="page_not_analysed", severity=Severity.WARNING, page_id=page.page_id,
                message=f"Page {page.page_number} non analysée : section vide dans le DOCX."))
            return

        anchored = [*page.backgrounds, *[image for image in page.image_elements if image.positioning != "flow"]]
        if anchored:
            carrier = self.add_carrier()
            for image in anchored:
                self.add_anchor(carrier, page, image)
        for element in page.text_elements:
            if not element.text.strip() and element.positioning in ("absolute", "flow"):
                self.report.issues.append(Issue(code="empty_text_block", severity=Severity.INFO,
                                                page_id=page.page_id, block_id=element.block_id,
                                                message="Bloc de texte vide ignoré."))
            elif element.positioning == "absolute":
                self.add_frame(page, element)
        self.add_flow(page, section)

    def add_flow(self, page: ComposedPage, first_section) -> None:
        bands = {band.index: band for band in page.bands}
        current_band: int | None = None
        current_column = 0
        started = False
        body_left = page.margins.left_pt
        body_width = page.width_pt - page.margins.left_pt - page.margins.right_pt
        for element in page.elements:
            if element.positioning != "flow":
                continue
            if isinstance(element, TextElement) and not element.text.strip():
                continue
            band_index = element.band_index if element.band_index is not None else (current_band or 0)
            column = element.column_index or 0
            column_break = False
            if not started or band_index != current_band:
                section = first_section if not started else self.document.add_section(WD_SECTION.CONTINUOUS)
                if started:
                    self.configure_section(section, page, WD_SECTION.CONTINUOUS)
                self.configure_band(section, page, bands.get(band_index))
                started, current_band, current_column = True, band_index, column
            elif column != current_column:
                columns = bands[band_index].columns if band_index in bands else 1
                column_break = columns > 1 and column > current_column
                current_column = column
            if isinstance(element, TextElement):
                self.add_text_paragraph(element, column_break)
                continue
            band = bands.get(band_index)
            if band is not None and band.column_rects and column < len(band.column_rects):
                column_rect = band.column_rects[column]
                left, width = column_rect.x_pt, column_rect.width_pt
            else:
                left, width = body_left, body_width
            self.add_inline_image(page, element, width, left, column_break)

    def report_fonts(self, pages: list[ComposedPage]) -> None:
        families = sorted({element.style.font_family for page in pages for element in page.text_elements})
        for family in families:
            used, fallback = fonts.resolve_family(family)
            if fallback:
                self.report.fallbacks.append(Issue(
                    code="font_not_installed", severity=Severity.INFO,
                    message=f"La police « {family} » n'est pas installée ici : le DOCX garde ce nom, Word ou "
                            f"LibreOffice la remplaceront si elle manque aussi chez le lecteur "
                            f"(équivalent local : « {used} »)."))


def _save_atomically(document, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        document.save(str(temporary))
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()


def _verify(output: Path, pages: list[ComposedPage], builder: _Builder, report: ExportReport) -> None:
    """Relit le fichier écrit : sections, textes, images, ancres."""
    inspection = inspect_docx(output)
    if inspection.page_sections != len(pages):
        report.issues.append(Issue(code="section_count_mismatch", severity=Severity.ERROR,
                                   message=f"Le DOCX relu a {inspection.page_sections} sections de page "
                                           f"au lieu de {len(pages)}."))
    corpus = "\n".join([inspection.text, *inspection.header_texts, *inspection.footer_texts])
    missing = [text for text in builder.expected_texts if text not in corpus]
    if missing:
        report.issues.append(Issue(code="text_missing", severity=Severity.ERROR,
                                   message=f"{len(missing)} bloc(s) de texte introuvable(s) dans le DOCX relu."))
    found_images = inspection.inline_images + inspection.anchored_images
    if found_images != builder.inserted_images:
        report.issues.append(Issue(code="image_count_mismatch", severity=Severity.ERROR,
                                   message=f"{found_images} image(s) relue(s) au lieu de {builder.inserted_images}."))
    for problem in inspection.anchor_problems:
        report.issues.append(Issue(code="anchor_invalid", severity=Severity.ERROR,
                                   message=f"Ancre d'image incomplète : {problem}"))


def build_docx(pages: list[ComposedPage], project: Project, output: Path) -> ExportReport:
    """Construit le DOCX (une section par page, mode lu dans chaque page) puis le relit pour le vérifier."""
    if not pages:
        raise ExportError("Aucune page à exporter.")
    output = Path(output)
    modes = sorted({str(page.mode) for page in pages})
    report = ExportReport(output=str(output), format="docx", mode=modes[0] if len(modes) == 1 else ",".join(modes))
    document = docx.Document()
    document.core_properties.title = project.name
    builder = _Builder(document, project, pages, report)
    builder.report_fonts(pages)
    for index, page in enumerate(pages):
        builder.add_page(page, first=index == 0)
    _save_atomically(document, output)
    report.page_count = len(pages)
    _verify(output, pages, builder, report)
    logger.info("DOCX écrit : %s (%d pages, %d blocs, %d images)", output, len(pages), report.text_blocks,
                report.images)
    return report


# --------------------------------------------------------------------------- relecture


class DocxInspection(BaseModel):
    """Contenu relu d'un DOCX (python-docx + XML)."""

    sections: int
    page_sections: int                          # sections commençant une nouvelle page (= pages source)
    section_sizes_pt: list[tuple[float, float]]
    paragraphs_by_style: dict[str, int]         # paragraphes non vides du corps (cadres compris)
    header_footer_styles: dict[str, int] = Field(default_factory=dict)
    non_empty_paragraphs: int
    frames: int
    inline_images: int
    anchored_images: int
    background_images: int = 0                  # ancres behindDoc="1"
    image_extents_pt: list[tuple[float, float]] = Field(default_factory=list)
    media_files: list[str]
    column_sections: int = 0                    # sections à plusieurs colonnes
    column_breaks: int = 0
    explicit_page_breaks: int = 0
    styles: list[str] = Field(default_factory=list)   # styles de paragraphe du livre présents
    header_texts: list[str] = Field(default_factory=list)
    footer_texts: list[str] = Field(default_factory=list)
    anchor_problems: list[str] = Field(default_factory=list)
    text: str


def _extent_pt(element: etree._Element) -> tuple[float, float]:
    extent = element.find("wp:extent", NS)
    return (round(int(extent.get("cx")) / EMU_PER_PT, 2), round(int(extent.get("cy")) / EMU_PER_PT, 2))


def inspect_docx(path: Path) -> DocxInspection:
    """Relecture réelle du fichier : ZIP, python-docx, XML du corps, en-têtes et pieds."""
    path = Path(path)
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ExportError("Archive DOCX corrompue.", detail=str(path))
            media = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
        document = docx.Document(str(path))
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as error:
        raise ExportError("DOCX illisible.", detail=f"{path}: {error}") from error

    body = document.element.body
    paragraphs = document.paragraphs
    styles = Counter(paragraph.style.name for paragraph in paragraphs if paragraph.text.strip())
    anchors = body.findall(".//wp:anchor", NS)
    inlines = body.findall(".//wp:inline", NS)
    header_texts, footer_texts = [], []
    header_footer_styles: Counter = Counter()
    for section in document.sections:
        for part, texts in ((section.header, header_texts), (section.footer, footer_texts)):
            if part.is_linked_to_previous:
                continue
            for paragraph in part.paragraphs:
                if paragraph.text.strip():
                    texts.append(paragraph.text)
                    header_footer_styles[paragraph.style.name] += 1
    problems = [f"ancre {index + 1} : {problem}" for index, anchor in enumerate(anchors)
                for problem in anchor_problems(anchor)]
    book_styles = set(STYLE_NAMES.values())
    return DocxInspection(
        sections=len(document.sections),
        page_sections=sum(1 for index, section in enumerate(document.sections)
                          if index == 0 or section.start_type != WD_SECTION.CONTINUOUS),
        section_sizes_pt=[(round(section.page_width.pt, 2), round(section.page_height.pt, 2))
                          for section in document.sections],
        paragraphs_by_style=dict(styles),
        header_footer_styles=dict(header_footer_styles),
        non_empty_paragraphs=sum(styles.values()),
        frames=len(body.findall(".//w:framePr", NS)),
        inline_images=len(inlines),
        anchored_images=len(anchors),
        background_images=sum(1 for anchor in anchors if anchor.get("behindDoc") == "1"),
        image_extents_pt=[_extent_pt(element) for element in [*inlines, *anchors]],
        media_files=media,
        column_sections=sum(1 for cols in body.iter(qn("w:cols")) if int(cols.get(qn("w:num"), "1")) > 1),
        column_breaks=len(body.findall(".//w:br[@w:type='column']", NS)),
        explicit_page_breaks=len(body.findall(".//w:br[@w:type='page']", NS)),
        styles=sorted(style.name for style in document.styles if style.name in book_styles),
        header_texts=header_texts,
        footer_texts=footer_texts,
        anchor_problems=problems,
        text="\n".join(paragraph.text for paragraph in paragraphs),
    )
