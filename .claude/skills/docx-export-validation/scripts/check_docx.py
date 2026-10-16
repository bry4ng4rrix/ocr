"""Contrôle d'intégrité d'un fichier DOCX (lecture seule).

Usage : python check_docx.py FICHIER.docx [--expect-pages N] [--expect-images N] [--expect-text MOT ...]
Code de sortie : 0 si tout est conforme, 1 si une attente échoue, 2 si le fichier est illisible.
"""
from __future__ import annotations

import argparse
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def inspect(path: Path) -> dict:
    report: dict = {"file": str(path), "size_bytes": path.stat().st_size}
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        names = archive.namelist()
    report["zip_ok"] = bad is None
    report["media_files"] = sorted(n for n in names if n.startswith("word/media/"))

    import docx  # python-docx

    document = docx.Document(str(path))
    body = document.element.body
    paragraphs = document.paragraphs
    styles = Counter(p.style.name for p in paragraphs if p.text.strip())
    fonts = Counter(r.font.name for p in paragraphs for r in p.runs if r.font.name)
    report.update({
        "sections": len(document.sections),
        "section_sizes_mm": [
            [round(s.page_width.mm, 1), round(s.page_height.mm, 1)] for s in document.sections
        ],
        "non_empty_paragraphs": sum(styles.values()),
        "paragraphs_by_style": dict(styles),
        "explicit_page_breaks": len(body.findall(".//w:br[@w:type='page']", NS)),
        "column_sections": len(body.findall(".//w:sectPr/w:cols[@w:num]", NS)),
        "frames": len(body.findall(".//w:framePr", NS)),
        "inline_images": len(body.findall(".//wp:inline", NS)),
        "anchored_images": len(body.findall(".//wp:anchor", NS)),
        "fonts": dict(fonts),
        "tables": len(document.tables),
    })
    text = "\n".join(p.text for p in paragraphs)
    text += "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    report["_text"] = text
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--expect-pages", type=int, help="sections attendues (mode page par page)")
    parser.add_argument("--expect-images", type=int)
    parser.add_argument("--expect-text", nargs="*", default=[])
    options = parser.parse_args()
    try:
        report = inspect(options.docx)
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as error:
        print(json.dumps({"file": str(options.docx), "error": f"{type(error).__name__}: {error}"}, ensure_ascii=False))
        return 2

    text = report.pop("_text")
    failures = []
    if not report["zip_ok"]:
        failures.append("archive ZIP corrompue")
    if options.expect_pages is not None and report["sections"] != options.expect_pages:
        failures.append(f"sections {report['sections']} ≠ {options.expect_pages} attendues")
    if options.expect_images is not None and len(report["media_files"]) < options.expect_images:
        failures.append(f"images {len(report['media_files'])} < {options.expect_images} attendues")
    missing = [t for t in options.expect_text if t not in text]
    if missing:
        failures.append(f"texte absent : {missing}")
    if report["non_empty_paragraphs"] == 0 and report["frames"] == 0:
        failures.append("aucun texte modifiable")
    report["failures"] = failures
    report["ok"] = not failures
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())
