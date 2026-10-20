"""Export DOCX : relecture réelle des fichiers produits (python-docx + XML).

Les pages composées sont construites ici (composition_service est développé en parallèle) en imitant
les règles de positionnement de D-06. Les constructeurs sont aussi utilisés par test_pdf_service.py.
"""
from __future__ import annotations

import hashlib
import json
import math
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

import docx
import numpy as np
import pytest
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

from models import (Alignment, BlockKind, ColumnBand, ComposedPage, ImageElement, Margins, PageCategory, Project,
                    Rect, ReconstructionMode, ResolvedStyle, Severity, TextElement)
from services import document_service
from tests.fixtures import synthetic
from services.document_service import (STYLE_NAMES, anchor_problems, build_docx, columns_element, frame_element,
                                       inspect_docx, style_differences)

ROOT = Path(__file__).resolve().parents[1]
CHECK_DOCX = ROOT / ".claude/skills/docx-export-validation/scripts/check_docx.py"

PAGE_W, PAGE_H = 419.53, 595.28          # A5 en points
MARGINS = Margins(top_pt=50, right_pt=45, bottom_pt=50, left_pt=45)
BODY_X = MARGINS.left_pt
BODY_W = PAGE_W - MARGINS.left_pt - MARGINS.right_pt
ALL_MODES = list(ReconstructionMode)

SERIF = "Liberation Serif"
SANS = "Liberation Sans"

STYLES: dict[BlockKind, ResolvedStyle] = {
    BlockKind.HEADING_1: ResolvedStyle(font_family=SERIF, font_size_pt=20, bold=True, alignment=Alignment.CENTER),
    BlockKind.HEADING_2: ResolvedStyle(font_family=SERIF, font_size_pt=14, bold=True),
    BlockKind.PARAGRAPH: ResolvedStyle(font_family=SERIF, font_size_pt=10.5, alignment=Alignment.JUSTIFY,
                                       first_line_indent_pt=14),
    BlockKind.LIST_ITEM: ResolvedStyle(font_family=SERIF, font_size_pt=10.5, left_indent_pt=14),
    BlockKind.CAPTION: ResolvedStyle(font_family=SERIF, font_size_pt=8.5, italic=True, alignment=Alignment.CENTER),
    BlockKind.HEADER: ResolvedStyle(font_family=SERIF, font_size_pt=8, italic=True, alignment=Alignment.CENTER),
    BlockKind.PAGE_NUMBER: ResolvedStyle(font_family=SERIF, font_size_pt=9, alignment=Alignment.CENTER),
}
WHITE_TITLE = ResolvedStyle(font_family=SANS, font_size_pt=26, bold=True, color="#FFFFFF",
                            alignment=Alignment.CENTER)
WHITE_TEXT = ResolvedStyle(font_family=SANS, font_size_pt=12, color="#FFFFFF", alignment=Alignment.CENTER)

PARAGRAPHS = [
    "Le lendemain matin, nous quittâmes la ville avant l'aube. La route longeait la rivière, "
    "et l'on apercevait déjà les premières maisons du village, serrées autour de l'église.",
    "Il faut avouer que le voyage fut plus long que prévu : les chevaux étaient fatigués, "
    "la pluie avait détrempé les chemins et notre guide hésitait à chaque carrefour.",
    "À midi, nous fîmes halte dans une auberge où l'on nous servit du pain, du fromage "
    "et un vin léger. Le propriétaire, un homme âgé, nous raconta l'histoire de la vallée.",
    "Le soir venu, la brume recouvrit les champs ; seules les cloches de l'abbaye "
    "rappelaient que nous n'étions pas seuls dans cette campagne silencieuse.",
]
LIST_ITEMS = ["– du pain de seigle et du fromage ;", "– une carte de la vallée ;", "– deux lanternes « tempête »."]


# --------------------------------------------------------------------------- constructeurs de pages composées


def estimated_height(text: str, style: ResolvedStyle, width: float) -> float:
    """Hauteur généreuse d'un bloc (largeur moyenne d'un caractère : 0,5 em)."""
    chars_per_line = max(1, int((width - style.left_indent_pt) / (0.5 * style.font_size_pt)))
    lines = sum(max(1, math.ceil((len(part) + 3) / chars_per_line)) for part in text.split("\n"))
    return round(lines * style.font_size_pt * style.line_spacing * 1.16 + 4, 1)


def positioning_for(mode: ReconstructionMode, kind: BlockKind, on_image: bool) -> str:
    """Règles de D-06 telles que les applique composition_service."""
    if mode == ReconstructionMode.FIDELITY:
        return "absolute"
    if kind == BlockKind.HEADER:
        return "header"
    if kind in (BlockKind.FOOTER, BlockKind.PAGE_NUMBER):
        return "footer"
    if mode == ReconstructionMode.HYBRID and on_image:
        return "absolute"
    return "flow"


class PageSketch:
    """Construit une ComposedPage réaliste : positions en points, ordre de lecture, écarts d'origine."""

    def __init__(self, page_id: str, number: int, mode: ReconstructionMode, category: PageCategory) -> None:
        self.page_id, self.number, self.mode, self.category = page_id, number, mode, category
        self.elements: list[TextElement | ImageElement] = []
        self.backgrounds: list[ImageElement] = []
        self.bands: list[ColumnBand] = []
        self.source_image: str | None = None
        self._count = 0

    def _id(self) -> str:
        self._count += 1
        return f"{self.page_id}-b{self._count}"

    def text(self, kind: BlockKind, text: str, y: float, x: float = BODY_X, width: float = BODY_W,
             style: ResolvedStyle | None = None, on_image: bool = False, band: int = 0, column: int = 0,
             height: float | None = None) -> float:
        style = style or STYLES[kind]
        height = height or estimated_height(text, style, width)
        self.elements.append(TextElement(
            block_id=self._id(), kind=kind, text=text, rect=Rect(x_pt=x, y_pt=y, width_pt=width, height_pt=height),
            style=style, positioning=positioning_for(self.mode, kind, on_image), on_image=on_image,
            band_index=band, column_index=column))
        return y + height

    def image(self, kind: str, file: Path, x: float, y: float, width: float, height: float, band: int = 0,
              column: int = 0) -> float:
        rect = Rect(x_pt=x, y_pt=y, width_pt=width, height_pt=height)
        if kind == "background":
            self.backgrounds.append(ImageElement(kind="background", file=str(file), rect=rect, z_index=0))
        else:
            positioning = "flow" if self.mode == ReconstructionMode.EDITABLE else "absolute"
            self.elements.append(ImageElement(block_id=self._id(), kind=kind, file=str(file), rect=rect,
                                              positioning=positioning, band_index=band, column_index=column))
        return y + height

    def build(self) -> ComposedPage:
        last_bottom: dict[tuple[int, int], float] = {}
        band_top = {band.index: band.top_pt for band in self.bands}
        for element in self.elements:
            if element.positioning != "flow":
                continue
            key = (element.band_index or 0, element.column_index or 0)
            previous = last_bottom.get(key, band_top.get(key[0], MARGINS.top_pt))
            element.gap_before_pt = round(max(0.0, element.rect.y_pt - previous), 1)
            last_bottom[key] = element.rect.bottom_pt
        return ComposedPage(page_id=self.page_id, page_number=self.number, width_pt=PAGE_W, height_pt=PAGE_H,
                            margins=MARGINS, mode=self.mode, category=self.category, backgrounds=self.backgrounds,
                            elements=self.elements, bands=self.bands, source_image=self.source_image)


def make_images(folder: Path) -> dict[str, Path]:
    """PNG de test : fond en dégradé texturé (cas B) et photographie bruitée (cas C)."""
    folder.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(7)
    height, width = 417, 420          # ratio du rect de fond du cas B (page A5, 70 % de la hauteur)
    ramp = np.linspace(0, 1, height)[:, None]
    background = np.zeros((height, width, 3), dtype=np.float64)
    background[..., 0] = 20 + 40 * ramp
    background[..., 1] = 90 + 80 * ramp
    background[..., 2] = 60 + 30 * ramp
    background += rng.normal(0, 6, background.shape)
    photo = np.zeros((300, 480, 3), dtype=np.float64)
    photo[..., 0] = np.linspace(60, 200, 480)[None, :]
    photo[..., 1] = np.linspace(120, 80, 300)[:, None]
    photo[..., 2] = 90
    photo += rng.normal(0, 25, photo.shape)
    paths = {"background": folder / "fond.png", "photo": folder / "photo.png"}
    Image.fromarray(np.clip(background, 0, 255).astype(np.uint8)).save(paths["background"])
    Image.fromarray(np.clip(photo, 0, 255).astype(np.uint8)).save(paths["photo"])
    return paths


def make_source_image(page: ComposedPage, path: Path, dpi: int = 100) -> Path:
    """« Photo corrigée » de la page : fond papier, images et texte dessinés aux rects (pour le fac-similé)."""
    scale = dpi / 72
    canvas = Image.new("RGB", (round(page.width_pt * scale), round(page.height_pt * scale)), (250, 248, 242))
    draw = ImageDraw.Draw(canvas)
    font_path = synthetic.font_path("serif")
    for image in [*page.backgrounds, *page.image_elements]:
        box = (round(image.rect.x_pt * scale), round(image.rect.y_pt * scale))
        size = (round(image.rect.width_pt * scale), round(image.rect.height_pt * scale))
        canvas.paste(Image.open(image.file).convert("RGB").resize(size), box)
    for element in page.text_elements:
        size = max(6, round(element.style.font_size_pt * scale))
        font = ImageFont.truetype(font_path, size) if Path(font_path).is_file() else ImageFont.load_default()
        words, lines, current = element.text.split(), [], ""
        for word in words:
            trial = f"{current} {word}".strip()
            if draw.textlength(trial, font=font) > element.rect.width_pt * scale and current:
                lines.append(current)
                current = word
            else:
                current = trial
        lines.append(current)
        y = element.rect.y_pt * scale
        fill = (255, 255, 255) if element.on_image else (30, 30, 30)
        for line in lines:
            draw.text((element.rect.x_pt * scale, y), line, font=font, fill=fill)
            y += size * 1.2
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(path)
    return path


def case_a(mode: ReconstructionMode, number: int = 1) -> ComposedPage:
    """Cas A : en-tête courant, titre, sous-titre, paragraphes, liste, numéro de page."""
    sketch = PageSketch(f"pA{number}", number, mode, PageCategory.TEXT)
    sketch.text(BlockKind.HEADER, "Voyage en Provence", y=24, height=12)
    y = sketch.text(BlockKind.HEADING_1, "Chapitre premier", y=80)
    y = sketch.text(BlockKind.HEADING_2, "I. Le départ", y=y + 14)
    y = sketch.text(BlockKind.PARAGRAPH, PARAGRAPHS[0], y=y + 10)
    y = sketch.text(BlockKind.PARAGRAPH, PARAGRAPHS[1], y=y + 4)
    for item in LIST_ITEMS:
        y = sketch.text(BlockKind.LIST_ITEM, item, y=y + 2)
    y = sketch.text(BlockKind.PARAGRAPH, PARAGRAPHS[2], y=y + 6)
    assert y < PAGE_H - MARGINS.bottom_pt, "le cas A doit tenir dans la page"
    sketch.text(BlockKind.PAGE_NUMBER, "12", y=560, x=190, width=40, height=14)
    return sketch.build()


def case_b(mode: ReconstructionMode, images: dict[str, Path], number: int = 2) -> ComposedPage:
    """Cas B : fond graphique sur 70 % de la page, texte clair posé dessus, texte ordinaire dessous."""
    sketch = PageSketch(f"pB{number}", number, mode, PageCategory.GRAPHIC_BACKGROUND)
    sketch.image("background", images["background"], 0, 0, PAGE_W, PAGE_H * 0.7)
    y = sketch.text(BlockKind.HEADING_1, "Les Couleurs du Sud", y=150, style=WHITE_TITLE, on_image=True)
    sketch.text(BlockKind.SPECIAL, "Lavande, olivier et pierre blonde", y=y + 20, style=WHITE_TEXT, on_image=True)
    y = sketch.text(BlockKind.PARAGRAPH, PARAGRAPHS[3], y=PAGE_H * 0.7 + 20)
    assert y < PAGE_H - MARGINS.bottom_pt
    return sketch.build()


def case_c(mode: ReconstructionMode, images: dict[str, Path], number: int = 3) -> ComposedPage:
    """Cas C : photographie, légende, trois paragraphes."""
    sketch = PageSketch(f"pC{number}", number, mode, PageCategory.PHOTO_TEXT)
    y = sketch.image("photo", images["photo"], BODY_X + 20, 60, BODY_W - 40, (BODY_W - 40) * 300 / 480)
    y = sketch.text(BlockKind.CAPTION, "Fig. 1 — Le village au lever du jour.", y=y + 6)
    for text in PARAGRAPHS[:3]:
        y = sketch.text(BlockKind.PARAGRAPH, text, y=y + 8)
    assert y < PAGE_H - MARGINS.bottom_pt
    return sketch.build()


def two_columns(mode: ReconstructionMode, number: int = 4) -> ComposedPage:
    """Titre pleine largeur (bande 0, une colonne) puis deux colonnes (bande 1)."""
    sketch = PageSketch(f"p2c{number}", number, mode, PageCategory.TEXT)
    gap = 15.0
    column_width = (BODY_W - gap) / 2
    left = Rect(x_pt=BODY_X, y_pt=120, width_pt=column_width, height_pt=420)
    right = Rect(x_pt=BODY_X + column_width + gap, y_pt=120, width_pt=column_width, height_pt=420)
    sketch.bands = [ColumnBand(index=0, top_pt=MARGINS.top_pt, bottom_pt=115, columns=1,
                               column_rects=[Rect(x_pt=BODY_X, y_pt=50, width_pt=BODY_W, height_pt=65)]),
                    ColumnBand(index=1, top_pt=120, bottom_pt=540, columns=2, gap_pt=gap,
                               column_rects=[left, right])]
    sketch.text(BlockKind.HEADING_1, "Deux colonnes", y=70, band=0)
    for column, rect in enumerate((left, right)):
        y = rect.y_pt
        for text in (PARAGRAPHS[column * 2], PARAGRAPHS[column * 2 + 1]):
            y = sketch.text(BlockKind.PARAGRAPH, text, y=y + 6, x=rect.x_pt, width=rect.width_pt, band=1,
                            column=column)
    return sketch.build()


def unanalysed(mode: ReconstructionMode, number: int = 5) -> ComposedPage:
    return ComposedPage(page_id=f"pX{number}", page_number=number, width_pt=PAGE_W, height_pt=PAGE_H,
                        margins=MARGINS, mode=mode, analysed=False)


def book(mode: ReconstructionMode, folder: Path, with_unanalysed: bool = True) -> list[ComposedPage]:
    """Livre de test : cas A, B, C, deux colonnes et (option) une page non analysée."""
    images = make_images(folder / "images")
    pages = [case_a(mode, 1), case_b(mode, images, 2), case_c(mode, images, 3), two_columns(mode, 4)]
    if with_unanalysed:
        pages.append(unanalysed(mode, 5))
    return pages


def project() -> Project:
    return Project(name="Livre de test")


def expected_image_count(pages: list[ComposedPage]) -> int:
    return sum(len(page.backgrounds) + len(page.image_elements) for page in pages if page.analysed)


# --------------------------------------------------------------------------- fixtures


@pytest.fixture(params=ALL_MODES, ids=[mode.value for mode in ALL_MODES])
def built(request, tmp_path):
    mode = request.param
    pages = book(mode, tmp_path)
    output = tmp_path / f"livre_{mode.value}.docx"
    report = build_docx(pages, project(), output)
    return mode, pages, output, report


# --------------------------------------------------------------------------- tests (exigences 14 à 18)


def test_docx_opens(built):
    mode, pages, output, report = built
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None
    document = docx.Document(str(output))
    assert len(document.paragraphs) > 0
    assert report.ok, [issue.message for issue in report.issues if issue.severity == Severity.ERROR]
    assert report.page_count == len(pages)
    assert report.mode == mode.value
    assert report.text_blocks == sum(len(page.text_elements) for page in pages)
    inspection = inspect_docx(output)
    assert inspection.anchor_problems == []


def test_text_is_real_paragraphs(built):
    """Chaque bloc est un paragraphe distinct (corps, cadre, en-tête ou pied), jamais une image."""
    _mode, pages, output, _report = built
    document = docx.Document(str(output))
    body_texts = [paragraph.text for paragraph in document.paragraphs]
    margin_texts = [paragraph.text for section in document.sections
                    for part in (section.header, section.footer) if not part.is_linked_to_previous
                    for paragraph in part.paragraphs]
    in_margins = Counter(e.text for page in pages for e in page.text_elements
                         if e.positioning in ("header", "footer"))
    in_body = Counter(e.text for page in pages for e in page.text_elements if e.positioning in ("flow", "absolute"))
    for text, count in in_body.items():
        assert body_texts.count(text) == count, text          # un paragraphe par bloc, texte entier
    for text, count in in_margins.items():
        assert margin_texts.count(text) == count, text
    inspection = inspect_docx(output)
    body_blocks = sum(1 for page in pages for e in page.text_elements if e.positioning in ("flow", "absolute"))
    assert inspection.non_empty_paragraphs == body_blocks
    assert "fromage" in inspection.text and "Les Couleurs du Sud" in inspection.text


def test_images_present(built):
    mode, pages, output, report = built
    inspection = inspect_docx(output)
    distinct_files = {image.file for page in pages for image in [*page.backgrounds, *page.image_elements]}
    assert len(inspection.media_files) == len(distinct_files)
    assert inspection.inline_images + inspection.anchored_images == expected_image_count(pages)
    assert report.images == expected_image_count(pages)
    assert inspection.background_images == sum(len(page.backgrounds) for page in pages)
    if mode == ReconstructionMode.EDITABLE:
        assert inspection.inline_images == 1          # la photo du cas C, dans le flux
    else:
        assert inspection.inline_images == 0          # tout est ancré


def test_editable_has_no_full_page_image(tmp_path):
    pages = book(ReconstructionMode.EDITABLE, tmp_path)
    for page in pages[:3]:
        page.source_image = str(make_source_image(page, tmp_path / "sources" / f"{page.page_id}.png"))
    output = tmp_path / "modifiable.docx"
    report = build_docx(pages, project(), output)
    inspection = inspect_docx(output)
    page_area = PAGE_W * PAGE_H
    assert all(width * height < 0.9 * page_area for width, height in inspection.image_extents_pt)
    with zipfile.ZipFile(output) as archive:
        media_hashes = {hashlib.sha256(archive.read(name)).hexdigest() for name in inspection.media_files}
    source_hashes = {hashlib.sha256(Path(page.source_image).read_bytes()).hexdigest() for page in pages[:3]}
    assert not media_hashes & source_hashes, "la photo de page ne doit jamais entrer dans le DOCX"
    assert not [issue for issue in report.issues if issue.code == "full_page_image_editable"]


def test_source_photo_is_refused_and_reported(tmp_path):
    images = make_images(tmp_path / "images")
    page = case_c(ReconstructionMode.FIDELITY, images)
    source = make_source_image(page, tmp_path / "source.png")
    page.source_image = str(source)
    page.elements.append(ImageElement(block_id="full", kind="photo", file=str(source),
                                      rect=Rect(x_pt=0, y_pt=0, width_pt=PAGE_W, height_pt=PAGE_H)))
    report = build_docx([page], project(), tmp_path / "refus.docx")
    assert any(issue.code == "source_photo_refused" and issue.severity == Severity.ERROR for issue in report.issues)
    assert inspect_docx(tmp_path / "refus.docx").anchored_images == 1      # seulement la vraie photo


def test_one_section_per_page(built):
    _mode, pages, output, _report = built
    inspection = inspect_docx(output)
    assert inspection.page_sections == len(pages)
    assert inspection.explicit_page_breaks == 0, "le changement de section suffit : pas de page blanche en plus"
    for width, height in inspection.section_sizes_pt:
        assert width == pytest.approx(PAGE_W, abs=0.1) and height == pytest.approx(PAGE_H, abs=0.1)
    document = docx.Document(str(output))
    first_of_pages = [section for index, section in enumerate(document.sections)
                      if index == 0 or section.start_type != WD_SECTION.CONTINUOUS]
    for section in first_of_pages:
        assert section.left_margin.pt == pytest.approx(MARGINS.left_pt, abs=0.1)
        assert section.top_margin.pt == pytest.approx(MARGINS.top_pt, abs=0.1)


def test_styles_are_reusable_and_used(built):
    mode, _pages, output, _report = built
    inspection = inspect_docx(output)
    assert set(inspection.styles) == set(STYLE_NAMES.values())
    for name in ("Titre 1", "Titre 2", "Corps", "Liste", "Légende"):
        assert inspection.paragraphs_by_style.get(name, 0) >= 1, name
    if mode != ReconstructionMode.FIDELITY:
        assert inspection.header_footer_styles.get("En-tête livre") == 1
        assert inspection.header_footer_styles.get("Numéro de page") == 1
    document = docx.Document(str(output))
    corps = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Corps"]
    # aucun « Corps » ne porte de taille ou de police en direct : un changement de style s'applique partout
    assert all(run.font.size is None and run.font.name is None for paragraph in corps for run in paragraph.runs)
    document.styles["Corps"].font.size = Pt(13)
    assert all(paragraph.style.font.size == Pt(13) for paragraph in corps)
    heading = document.styles["Titre 1"].element
    assert heading.find(f".//{qn('w:outlineLvl')}").get(qn("w:val")) == "0"


def test_direct_formatting_only_for_differences(tmp_path):
    page = case_a(ReconstructionMode.EDITABLE)
    special = next(e for e in page.text_elements if e.kind == BlockKind.PARAGRAPH)
    special.style = special.style.model_copy(update={"font_size_pt": 12.5, "italic": True})
    output = tmp_path / "surcharge.docx"
    build_docx([page], project(), output)
    document = docx.Document(str(output))
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    run = paragraphs[special.text].runs[0]
    assert run.font.size == Pt(12.5) and run.font.italic is True
    others = [p for p in document.paragraphs if p.style.name == "Corps" and p.text and p.text != special.text]
    assert others and all(r.font.size is None and r.font.italic is None for p in others for r in p.runs)


def test_style_differences_tolerances():
    base = STYLES[BlockKind.PARAGRAPH]
    assert style_differences(base.model_copy(update={"font_size_pt": 10.6}), base) == {}
    differences = style_differences(base.model_copy(update={"font_size_pt": 12, "color": "#FF0000",
                                                            "space_before_pt": 9}), base)
    assert differences == {"font_size_pt": 12, "color": "#FF0000", "space_before_pt": 9}


def test_columns_in_flow(tmp_path):
    for mode in (ReconstructionMode.EDITABLE, ReconstructionMode.HYBRID):
        page = two_columns(mode)
        output = tmp_path / f"colonnes_{mode.value}.docx"
        report = build_docx([page], project(), output)
        assert report.ok
        inspection = inspect_docx(output)
        assert inspection.page_sections == 1 and inspection.sections == 2   # bande 0 puis section continue
        assert inspection.column_sections == 1
        assert inspection.column_breaks == 1
        document = docx.Document(str(output))
        cols = document.sections[1]._sectPr.find(qn("w:cols"))
        assert cols.get(qn("w:num")) == "2" and cols.get(qn("w:space")) == "300"
        assert document.sections[1].start_type == WD_SECTION.CONTINUOUS
        broken = [p for p in document.paragraphs if p._p.findall(f".//{qn('w:br')}")]
        assert len(broken) == 1 and broken[0].text.strip() == PARAGRAPHS[2]


def test_flow_spacing_uses_original_gap(tmp_path):
    page = case_a(ReconstructionMode.EDITABLE)
    output = tmp_path / "ecarts.docx"
    build_docx([page], project(), output)
    document = docx.Document(str(output))
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}
    for element in page.text_elements:
        if element.positioning != "flow":
            continue
        paragraph = by_text[element.text]
        space_before = paragraph.paragraph_format.space_before
        if space_before is None:
            space_before = paragraph.style.paragraph_format.space_before
        assert space_before.pt == pytest.approx(element.gap_before_pt, abs=0.8), element.text


def test_frames_in_fidelity(tmp_path):
    page = case_a(ReconstructionMode.FIDELITY)
    output = tmp_path / "fidelite.docx"
    build_docx([page], project(), output)
    inspection = inspect_docx(output)
    assert inspection.frames == len(page.text_elements)
    document = docx.Document(str(output))
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}
    for element in page.text_elements:
        frame = by_text[element.text]._p.find(f".//{qn('w:framePr')}")
        assert frame is not None
        assert int(frame.get(qn("w:x"))) == round(element.rect.x_pt * 20)
        assert int(frame.get(qn("w:y"))) == round(element.rect.y_pt * 20)
        assert int(frame.get(qn("w:w"))) == round(element.rect.width_pt * 20)
        assert int(frame.get(qn("w:h"))) == round(element.rect.height_pt * 20)
        assert frame.get(qn("w:hAnchor")) == "page" and frame.get(qn("w:vAnchor")) == "page"
        assert frame.get(qn("w:hRule")) == "atLeast"


def test_hybrid_frames_only_for_text_on_image(tmp_path):
    images = make_images(tmp_path / "images")
    pages = [case_b(ReconstructionMode.HYBRID, images), case_c(ReconstructionMode.HYBRID, images)]
    output = tmp_path / "hybride.docx"
    build_docx(pages, project(), output)
    inspection = inspect_docx(output)
    assert inspection.frames == sum(1 for page in pages for e in page.text_elements if e.on_image) == 2
    body = docx.Document(str(output)).element.body
    anchors = body.findall(f".//{qn('wp:anchor')}")
    background = [a for a in anchors if a.get("behindDoc") == "1"]
    photo = [a for a in anchors if a.get("behindDoc") == "0"]
    assert len(background) == 1 and background[0].find(qn("wp:wrapNone")) is not None
    assert len(photo) == 1 and photo[0].find(qn("wp:wrapSquare")) is not None
    position = photo[0].find(qn("wp:positionV"))
    assert position.get("relativeFrom") == "page"
    assert int(position.find(qn("wp:posOffset")).text) == round(60 * 12700)


def test_header_and_footer(tmp_path):
    pages = [case_a(ReconstructionMode.EDITABLE, 1), case_a(ReconstructionMode.EDITABLE, 2),
             case_c(ReconstructionMode.EDITABLE, make_images(tmp_path / "images"), 3)]
    output = tmp_path / "entetes.docx"
    build_docx(pages, project(), output)
    document = docx.Document(str(output))
    sections = document.sections
    assert sections[0].header.paragraphs[0].text == "Voyage en Provence"
    assert sections[0].footer.paragraphs[0].text == "12"
    assert sections[1].header.is_linked_to_previous          # même en-tête : section liée
    assert not sections[2].header.is_linked_to_previous      # page sans en-tête : en-tête vide propre
    assert all(not paragraph.text for paragraph in sections[2].header.paragraphs)
    assert sections[0].header_distance.pt == pytest.approx(24, abs=0.1)
    assert sections[0].footer_distance.pt == pytest.approx(PAGE_H - 574, abs=0.1)


def test_unanalysed_page_gives_empty_section_and_issue(tmp_path):
    pages = [case_a(ReconstructionMode.EDITABLE, 1), unanalysed(ReconstructionMode.EDITABLE, 2)]
    output = tmp_path / "non_analysee.docx"
    report = build_docx(pages, project(), output)
    assert any(issue.code == "page_not_analysed" and issue.page_id == "pX2" for issue in report.issues)
    assert report.ok
    inspection = inspect_docx(output)
    assert inspection.page_sections == 2
    assert inspection.inline_images + inspection.anchored_images == 0


def test_missing_image_is_reported_and_text_kept(tmp_path):
    images = make_images(tmp_path / "images")
    page = case_c(ReconstructionMode.HYBRID, images)
    page.image_elements[0].file = str(tmp_path / "absente.png")
    report = build_docx([page], project(), tmp_path / "manque.docx")
    assert not report.ok and any(issue.code == "image_missing" for issue in report.issues)
    assert PARAGRAPHS[0] in inspect_docx(tmp_path / "manque.docx").text


def test_font_not_installed_is_listed_as_fallback(tmp_path):
    page = case_a(ReconstructionMode.EDITABLE)
    page.text_elements[1].style = page.text_elements[1].style.model_copy(update={"font_family": "Garamond Imaginaire"})
    output = tmp_path / "police.docx"
    report = build_docx([page], project(), output)
    assert any(issue.code == "font_not_installed" and "Garamond Imaginaire" in issue.message
               for issue in report.fallbacks)
    document = docx.Document(str(output))
    # seul titre de niveau 1 du livre : sa police devient la base du style « Titre 1 » (pas de mise en forme directe)
    run = next(p for p in document.paragraphs if p.text == "Chapitre premier").runs[0]
    assert run.font.name is None
    fonts_element = document.styles["Titre 1"].element.rPr.find(qn("w:rFonts"))
    assert {fonts_element.get(qn(f"w:{a}")) for a in ("ascii", "hAnsi", "eastAsia", "cs")} == {"Garamond Imaginaire"}
    # un sous-titre isolé dans une autre police porte, lui, la police en direct
    page.text_elements[2].style = page.text_elements[2].style.model_copy(update={"font_family": "Autre Police"})
    build_docx([page, case_a(ReconstructionMode.EDITABLE, 2), case_a(ReconstructionMode.EDITABLE, 3)], project(),
               output)
    document = docx.Document(str(output))
    runs = [p.runs[0] for p in document.paragraphs if p.text == "I. Le départ"]
    assert sorted(str(run.font.name) for run in runs) == ["Autre Police", "None", "None"]
    direct = next(run for run in runs if run.font.name == "Autre Police")._r.rPr.find(qn("w:rFonts"))
    assert {direct.get(qn(f"w:{a}")) for a in ("ascii", "hAnsi", "eastAsia", "cs")} == {"Autre Police"}


def test_atomic_save_replaces_existing_file(tmp_path):
    output = tmp_path / "livre.docx"
    output.write_bytes(b"ancienne version")
    build_docx([case_a(ReconstructionMode.EDITABLE)], project(), output)
    assert zipfile.is_zipfile(output)
    assert [path.name for path in tmp_path.iterdir() if path.suffix == ".tmp"] == []


def test_empty_book_is_refused(tmp_path):
    from models.errors import ExportError
    with pytest.raises(ExportError):
        build_docx([], project(), tmp_path / "vide.docx")


# --------------------------------------------------------------------------- fragments XML


def test_frame_element_values():
    frame = frame_element(Rect(x_pt=10, y_pt=20.5, width_pt=100, height_pt=30), wrap="around")
    assert frame.get(qn("w:x")) == "200" and frame.get(qn("w:y")) == "410"
    assert frame.get(qn("w:w")) == "2000" and frame.get(qn("w:h")) == "600"
    assert frame.get(qn("w:wrap")) == "around" and frame.get(qn("w:hRule")) == "atLeast"


def test_columns_element_equal_and_unequal():
    equal = columns_element(2, 15)
    assert equal.get(qn("w:num")) == "2" and equal.get(qn("w:space")) == "300" and len(equal) == 0
    unequal = columns_element(2, 12, [200, 100])
    assert unequal.get(qn("w:equalWidth")) == "0"
    assert [col.get(qn("w:w")) for col in unequal] == ["4000", "2000"]


def test_anchor_problems_detects_incomplete_anchor(tmp_path):
    images = make_images(tmp_path / "images")
    output = tmp_path / "ancre.docx"
    build_docx([case_c(ReconstructionMode.FIDELITY, images)], project(), output)
    anchor = docx.Document(str(output)).element.body.find(f".//{qn('wp:anchor')}")
    assert anchor_problems(anchor) == []
    del anchor.attrib["relativeHeight"]
    anchor.remove(anchor.find(qn("wp:simplePos")))
    problems = anchor_problems(anchor)
    assert "attribut relativeHeight absent" in problems and "enfant simplePos absent" in problems


def test_section_xml_children_order(built):
    """Les enfants de w:sectPr et w:pPr respectent l'ordre du schéma (sinon Word refuse le fichier)."""
    _mode, _pages, output, _report = built
    body = docx.Document(str(output)).element.body
    for sect_pr in body.iter(qn("w:sectPr")):
        names = [child.tag.split("}")[1] for child in sect_pr]
        ranks = [document_service.SECTPR_ORDER.index(name) for name in names]
        assert ranks == sorted(ranks), names
    for p_pr in body.iter(qn("w:pPr")):
        names = [child.tag.split("}")[1] for child in p_pr]
        ranks = [document_service.PPR_ORDER.index(name) for name in names]
        assert ranks == sorted(ranks), names


# --------------------------------------------------------------------------- script du skill


def test_check_docx_script_accepts_export(tmp_path):
    images = make_images(tmp_path / "images")
    pages = [case_a(ReconstructionMode.HYBRID), case_b(ReconstructionMode.HYBRID, images),
             case_c(ReconstructionMode.HYBRID, images)]
    output = tmp_path / "skill.docx"
    build_docx(pages, project(), output)
    done = subprocess.run([sys.executable, str(CHECK_DOCX), str(output), "--expect-pages", "3",
                           "--expect-images", "2", "--expect-text", "Chapitre premier", "Les Couleurs du Sud",
                           "Fig. 1"], capture_output=True, text=True, timeout=120)
    result = json.loads(done.stdout)
    assert done.returncode == 0, result["failures"]
    assert result["anchored_images"] == 2 and result["frames"] == 2
